from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


def write_pdf(path: Path, text: str = "Abstract. Research paper. References."):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 50 780 Td ({text}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as handle:
        writer.write(handle)


def make_examples(source: Path):
    source.mkdir(parents=True, exist_ok=True)
    (source / "实验报告_电路.txt").write_text("实验目的：测量电阻。\n实验步骤：连接电路。\n实验结果：符合预期。", encoding="utf-8")
    (source / "a").mkdir()
    (source / "b").mkdir()
    (source / "a" / "课程资料.txt").write_text("课程讲义：计算机网络第一章。", encoding="utf-8")
    (source / "b" / "课程资料.txt").write_text("课程讲义：计算机网络第二章。", encoding="gb18030")
    doc = Document()
    doc.add_heading("课程资料：数据库", 0)
    doc.add_paragraph("课程讲义与复习重点。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "课程"
    table.cell(0, 1).text = "索引与事务"
    doc.save(source / "数据库讲义.docx")
    write_pdf(source / "研究文献.pdf")
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with (source / "无正文.pdf").open("wb") as handle:
        writer.write(handle)
    (source / "损坏.pdf").write_bytes(b"this is deliberately not a PDF")
    (source / "损坏.docx").write_bytes(b"this is deliberately not a DOCX")
    (source / "个人笔记.txt").write_text("备忘录：明天带书。", encoding="utf-16")
    (source / "未知.txt").write_text("12345", encoding="utf-8")
    (source / "忽略.png").write_bytes(b"fixture")
    return source
