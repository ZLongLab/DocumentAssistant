"""Executed in a separate process so a broken parser cannot block the scan."""
from __future__ import annotations

import json
import sys
from pathlib import Path

TEXT_LIMIT = 100_000
BYTE_LIMIT = 50 * 1024 * 1024


def extract(path: Path) -> dict:
    try:
        if path.stat().st_size == 0:
            return {'text':'', 'status':'无法提取：空文件（0 字节）'}
        detected_docx = False
        metadata_title = ''
        if path.stat().st_size > BYTE_LIMIT:
            return {"text": "", "status": "未提取：文件超过 50 MB"}
        truncated = False
        if path.suffix.lower() == ".txt":
            raw = path.read_bytes()
            for encoding in ("utf-8-sig", "utf-16", "gb18030"):
                if encoding == "utf-16" and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
                    continue
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeError:
                    pass
            else:
                return {"text": "", "status": "无法提取：不支持的 TXT 编码"}
        elif path.suffix.lower() == ".docx":
            from docx import Document
            doc = Document(path)
            metadata_title = doc.core_properties.title or ''
            parts = [p.text for p in doc.paragraphs]
            parts += [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
            text = "\n".join(parts)
        elif path.suffix.lower() == ".doc":
            import zipfile
            if zipfile.is_zipfile(path):
                from docx import Document
                doc = Document(path)
                metadata_title = doc.core_properties.title or ''
                text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
                detected_docx = True
            else:
                from .legacy_doc import extract_doc
                text, truncated = extract_doc(path, TEXT_LIMIT, BYTE_LIMIT)
        elif path.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(""):
                return {"text": "", "status": "无法提取：PDF 需要密码"}
            metadata_title = (reader.metadata.title if reader.metadata else '') or ''
            parts, count = [], 0
            for index, page in enumerate(reader.pages):
                chunk = page.extract_text() or ""
                parts.append(chunk)
                count += len(chunk)
                if count >= TEXT_LIMIT:
                    truncated = index < len(reader.pages) - 1
                    break
            text = "\n".join(parts)
        else:
            return {"text": "", "status": "不支持的文件类型"}
        truncated = truncated or len(text) > TEXT_LIMIT
        text = text[:TEXT_LIMIT].strip()
        if not text:
            return {"text": "", "status": "无可提取正文（可能是扫描件；本版无 OCR）"}
        from .document_metadata import title_suggestion
        title, title_source = title_suggestion(path.name,text,metadata_title)
        return {"display_title":title,"title_source":title_source,"text": text, "status": "已提取（前 10 万字，已截断）" if truncated else "已提取（DOC 实为 DOCX，已自动识别）" if detected_docx else "已提取"}
    except Exception as exc:
        return {"text": "", "status": f"无法提取：{type(exc).__name__}: {str(exc)[:180]}"}


if __name__ == "__main__":
    print(json.dumps(extract(Path(sys.argv[1])), ensure_ascii=True))
